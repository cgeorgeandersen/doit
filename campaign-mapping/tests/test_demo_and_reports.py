import csv
import io

import openpyxl
import pandas as pd
import pytest

from campaign_mapping import reports, taxonomy
from campaign_mapping.db import connect
from campaign_mapping.demo import DATA_DIR, ensure_demo_db
from campaign_mapping.engine import classify
from campaign_mapping.models import CLASSIFIED, NETWORK, WORKSPACE
from campaign_mapping.rulesets import list_versions, load_ruleset
from campaign_mapping.runs import existing_hashes, list_runs, load_records, prepare_rows, reproduce_run
from campaign_mapping.suggest import comparable, initials, suggest


@pytest.fixture(scope="module")
def frames(demo_db):
    conn = connect(demo_db)
    ruleset = load_ruleset(conn)
    result = classify(load_records(conn), ruleset)
    records = reports.records_frame(conn)
    outcomes = reports.outcomes_frame(result, ruleset)
    campaigns = pd.DataFrame([dict(row) for row in taxonomy.list_campaigns(conn)])
    yield conn, ruleset, result, records, outcomes, campaigns
    conn.close()


@pytest.fixture(scope="module")
def network(frames):
    _, ruleset, _, records, outcomes, _ = frames
    return reports.lens_view(records, outcomes, ruleset.dimensions)


def truth_by_source_row():
    with (DATA_DIR / "ground_truth.csv").open() as handle:
        return {int(row["row_id"]): row for row in csv.DictReader(handle)}


# ── the demo database ───────────────────────────────────────────────────────

def test_the_starter_ruleset_covers_sixty_to_seventy_percent_of_spend(frames, network):
    _, ruleset, *_ = frames
    campaign = reports.coverage_by_dimension(network, ruleset.dimensions).loc["campaign"]
    assert 0.60 <= campaign["spend_pct"] <= 0.70
    assert campaign["conflict_rows"] > 0            # the review queue has real work, conflicts included


def test_the_starter_rules_are_right_where_they_fire(network):
    truth = truth_by_source_row()
    fired = network[(network["dimension"] == "campaign") & (network["status"] == CLASSIFIED)]
    right = fired[[value == truth[row]["true_campaign"] for value, row in zip(fired["value"], fired["source_row"], strict=True)]]
    assert right["spend"].sum() / fired["spend"].sum() > 0.98


def test_the_demo_has_its_history_and_a_reproducible_starter_run(frames):
    conn = frames[0]
    assert [v.message for v in list_versions(conn)][-1].startswith("Starter ruleset")
    (run,) = list_runs(conn)
    assert reproduce_run(conn, run.id).matches


def test_the_ground_truth_never_reaches_the_app(frames):
    conn = frames[0]
    columns = {row[1] for table in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'")
               for row in conn.execute(f"PRAGMA table_info({table[0]})")}
    assert not any(column.startswith("true_") for column in columns)
    with (DATA_DIR / "utm_strings.csv").open() as handle:
        assert conn.execute("SELECT COUNT(*) FROM raw_strings").fetchone()[0] == sum(1 for _ in handle) - 1


def test_the_demo_database_is_built_once(tmp_path):
    path = ensure_demo_db(tmp_path / "demo.db")
    stamp = path.stat().st_mtime_ns
    assert ensure_demo_db(path).stat().st_mtime_ns == stamp


# ── coverage ────────────────────────────────────────────────────────────────

def test_every_dollar_is_accounted_for_once_per_dimension(frames, network):
    _, ruleset, _, records, *_ = frames
    coverage = reports.coverage_by_dimension(network, ruleset.dimensions)
    total = records["spend"].sum()
    for _, row in coverage.iterrows():
        parts = row[["classified_spend", "conflict_spend", "ignored_spend", "unclassified_spend"]].sum()
        assert parts == pytest.approx(total)
        assert row["spend_pct"] == pytest.approx(row["classified_spend"] / total)


def test_each_operator_sees_its_own_product_lineup(frames):
    _, ruleset, result, records, outcomes, _ = frames
    northgate = set(records[records["operator_slug"] == "northgate"]["raw_string_id"])
    string = next(
        c.raw_string_id
        for c in result.where(dimension="product", lens=NETWORK)
        if c.raw_string_id in northgate and c.method == "campaign_bridge"
        and ruleset.label(result.get(c.raw_string_id, "campaign").value_id) == "Summer Cup 2026"
    )
    network_values = set(ruleset.labels(result.get(string, "product", NETWORK).value_ids))
    own_values = set(ruleset.labels(result.get(string, "product", WORKSPACE).value_ids))
    assert own_values == {"Zestify Lime", "Zestify Original"}             # Northgate's own attribution
    assert network_values == {"Zestify Original", "Zestify Zero", "Zestify Lime"}   # the shared lineup
    by_operator = reports.coverage_by_operator(reports.lens_view(records, outcomes, ruleset.dimensions, reports.OWN_LENS),
                                               ruleset.dimensions)
    assert set(by_operator.index) == set(records["operator"])


def test_the_review_queue_is_ranked_by_spend_and_adds_up(frames, network):
    queue = reports.review_queue(network, "campaign", "campaign")
    assert list(queue["spend"]) == sorted(queue["spend"], reverse=True)
    assert queue["key"].is_unique
    assert queue.iloc[0]["key"] == "sc-26"
    waiting = network[(network["dimension"] == "campaign") & network["status"].isin(["unclassified", "conflict"])]
    assert queue["spend"].sum() == pytest.approx(waiting["spend"].sum())
    assert set(queue[queue["status"] == "conflict"]["key"]) >= {"summer_cup_to_fall_kickoff_bridge"}


def test_the_top_unrecognized_strings_are_unclassified_and_ranked(network):
    top = reports.unrecognized(network, "campaign", limit=20)
    assert len(top) == 20 and list(top["rank"]) == list(range(1, 21))
    assert list(top["spend"]) == sorted(top["spend"], reverse=True)


def test_rule_reach_previews_a_rule_before_it_exists(frames, network):
    _, _, _, records, *_ = frames
    totals, sample = reports.rule_reach(records, network, "campaign", field="campaign", match_type="exact", pattern="SC-26")
    queue = reports.review_queue(network, "campaign", "campaign")
    assert totals["rows"] == queue.iloc[0]["rows"] == totals["open_rows"]
    assert set(sample["status"]) == {"unclassified"}


# ── suggestions ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize(
    "text, expected",
    [
        ("sc-26", "Summer Cup 2026"),               # like classified "sc26_promo"
        ("sr26_promo", "Spring Refresh 2026"),      # "sr" spells its initials
        ("ww-26", "Winter Warmup 2026"),
        ("witer_warmup 2026", "Winter Warmup 2026"),   # a typo
        ("ff26_q3", "Pinecrest Fall Fest 2026"),    # initials of its last two words
        ("fal kickoff 2026", "Fall Kickoff 2026"),
    ],
)
def test_suggestions_find_what_a_string_means(frames, network, text, expected):
    _, ruleset, *_ = frames
    examples = reports.classified_examples(network, "campaign", "campaign")
    suggestions = suggest(text, "campaign", ruleset, examples)
    assert suggestions[0].value == expected
    assert suggestions[0].reason


def test_suggestions_ignore_numbers_and_know_initials():
    assert comparable("SC26_Promo") == comparable("sc promo 2026") == "sc promo"
    assert initials("Pinecrest Fall Fest 2026") == {"pff", "pf", "ff"}


# ── exports ─────────────────────────────────────────────────────────────────

def test_the_mapping_table_has_one_row_per_string_and_names_its_version(frames):
    _, ruleset, result, records, outcomes, _ = frames
    mapping = reports.mapping_table(records, outcomes, ruleset.dimensions, version_id=ruleset.version_id,
                                    engine_version=result.engine_version)
    assert len(mapping) == len(records)
    assert set(mapping["ruleset_version"]) == {ruleset.version_id}
    for column in ("campaign", "campaign_rule", "product", "product_operator_lens", "objective", "type", "utm_campaign"):
        assert column in mapping.columns
    assert mapping["campaign"].isin(["[unclassified]", "[conflict]"]).any()   # never a silent default


def test_the_campaign_rollup_adds_up_to_all_the_spend(frames, network):
    _, ruleset, _, records, _, campaigns = frames
    rollup = reports.campaign_rollup(network, ruleset, campaigns)
    assert rollup["spend"].sum() == pytest.approx(records["spend"].sum())
    assert rollup.iloc[0]["campaign"] == "Summer Cup 2026"
    assert "(unclassified)" in set(rollup["campaign"])
    assert "Zestify Lime" in rollup.set_index("campaign").loc["Summer Cup 2026", "network_products"]


def test_the_excel_export_has_every_sheet(frames, network):
    _, ruleset, _, _, _, campaigns = frames
    rollup = reports.campaign_rollup(network, ruleset, campaigns)
    book = openpyxl.load_workbook(io.BytesIO(reports.excel_bytes({"Mapping": rollup, "Campaign roll-up": rollup})))
    assert book.sheetnames == ["Mapping", "Campaign roll-up"]
    assert book["Mapping"].max_row == len(rollup) + 1


# ── import preview ──────────────────────────────────────────────────────────

def test_an_import_preview_spots_rows_already_loaded(frames):
    conn = frames[0]
    with (DATA_DIR / "utm_strings.csv").open() as handle:
        rows = [{**row, "source_row": row.pop("row_id")} for _, row in zip(range(50), csv.DictReader(handle), strict=False)]
    prepared = prepare_rows(conn, rows + [{**rows[0], "spend": "999999"}])
    already = existing_hashes(conn, [p.row_hash for p in prepared])
    assert sum(p.row_hash in already for p in prepared) == 50   # the 50 real rows, not the edited one
